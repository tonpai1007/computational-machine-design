"""Tests for the mdie command line interface."""

from __future__ import annotations

from pathlib import Path

import pytest
from rich.console import Console

from cli import main
from cli.app import SUBCOMMANDS, _design_flags, _resolve_target
from cli.interactive import render_header, render_menu_table, run_interactive_menu
from cli.viewers import find_cad_tools, launch_cad_viewer

MD_SAMPLE = "# Sheet\n\nbody\n"


def test_bare_invocation_opens_the_interactive_prompt(monkeypatch):
    called = []
    monkeypatch.setattr("cli.app.interactive_repl", lambda: called.append(True))
    assert main([]) == 0
    assert called == [True]


@pytest.mark.parametrize("flag", ["-h", "--help", "help"])
def test_help_exits_cleanly(flag, capsys):
    assert main([flag]) == 0
    out = capsys.readouterr().out
    assert "MDIE" in out
    assert "convert" in out


def test_info_lists_convertible_formats(capsys):
    assert main(["info"]) == 0
    out = capsys.readouterr().out
    assert "convert formats" in out
    assert "pdf" in out


def test_convert_list_formats_prints_markdown(capsys):
    assert main(["convert", "--list-formats"]) == 0
    assert "Drawing sheets" in capsys.readouterr().out


def test_convert_round_trip_through_the_cli(tmp_path, capsys):
    src = tmp_path / "r.md"
    src.write_text(MD_SAMPLE, encoding="utf-8")
    dst = tmp_path / "r.docx"

    assert main(["convert", str(src), "--to", "docx", "-o", str(dst)]) == 0
    assert dst.exists()
    assert "Wrote" in capsys.readouterr().out


def test_convert_reports_failure_without_a_traceback(tmp_path, capsys):
    src = tmp_path / "r.md"
    src.write_text(MD_SAMPLE, encoding="utf-8")
    assert main(["convert", str(src), "--to", "dwg", "-o", str(tmp_path / "r.dwg")]) == 1
    assert "Conversion failed" in capsys.readouterr().out


def test_legacy_docx_alias_still_works(tmp_path):
    src = tmp_path / "r.html"
    src.write_text("<html><body><h1>Hi</h1></body></html>", encoding="utf-8")
    dst = tmp_path / "r.docx"
    assert main(["docx", str(src), "-o", str(dst)]) == 0
    assert dst.exists()


def test_view_reports_a_missing_target(tmp_path):
    assert main(["view", str(tmp_path / "does_not_exist")]) == 1


def test_design_flags_are_extracted_in_any_order():
    assert _design_flags(["--open", "-o", "Project/x"]) == ("Project/x", True)
    assert _design_flags(["--output-dir", "Project/y"]) == ("Project/y", False)
    assert _design_flags([]) == (None, False)


def test_resolve_target_falls_back_to_the_project_folder(tmp_path, monkeypatch):
    """A bare project name resolves to Project/<name> when that folder exists.

    Built in a temp tree rather than against the real Project/chair, which is
    gitignored and therefore absent on a fresh CI checkout.
    """
    monkeypatch.chdir(tmp_path)
    (tmp_path / "Project" / "widget").mkdir(parents=True)

    assert _resolve_target("widget") == Path("Project") / "widget"


def test_resolve_target_returns_a_direct_file_path(tmp_path, monkeypatch):
    """A path that already exists is returned unchanged, with no fallback."""
    monkeypatch.chdir(tmp_path)
    part = tmp_path / "bracket.step"
    part.write_text("ISO-10303-21;", encoding="utf-8")

    assert _resolve_target("bracket.step") == Path("bracket.step")


def test_resolve_target_returns_the_input_when_nothing_exists(tmp_path, monkeypatch):
    """An unresolvable name comes back as-is so callers can report it."""
    monkeypatch.chdir(tmp_path)

    assert _resolve_target("nonexistent") == Path("nonexistent")


def test_known_subcommands_are_declared():
    for name in ("design", "convert", "view", "drawings", "report", "info"):
        assert name in SUBCOMMANDS


def test_legacy_subcommand_aliases_are_preserved():
    for alias in ("docx", "word", "convert-docx", "consolidate", "study", "solve", "optimize"):
        assert alias in SUBCOMMANDS


def test_process_prompt_reads_a_docx_spec_file(tmp_path):
    """A .docx path is extracted into text by the prompt reader."""
    import docx

    from cli.app import _maybe_read_file_prompt

    doc = docx.Document()
    doc.add_heading("Spec", level=1)
    doc.add_paragraph("Design a shaft 500 mm long with a keyway")
    spec = tmp_path / "spec.docx"
    doc.save(str(spec))

    text = _maybe_read_file_prompt(str(spec))
    assert "Design a shaft 500 mm long with a keyway" in text
    assert "Spec" in text


def test_process_prompt_reads_an_md_spec_file(tmp_path):
    from cli.app import _maybe_read_file_prompt

    spec = tmp_path / "spec.md"
    spec.write_text("# Bracket\nmotor mount, 4 bolt holes, 10 kg", encoding="utf-8")

    text = _maybe_read_file_prompt(str(spec))
    assert "motor mount" in text
    assert "4 bolt holes" in text


def test_process_prompt_passes_through_a_non_file_prompt():
    """A plain sentence is not treated as a document to read."""
    from cli.app import _maybe_read_file_prompt

    assert _maybe_read_file_prompt("design a chair") == "design a chair"


def test_find_cad_tools_returns_a_mapping():
    tools = find_cad_tools()
    assert isinstance(tools, dict)
    assert set(tools) <= {"openscad", "freecad"}


def test_launching_a_viewer_for_a_missing_file_returns_false(tmp_path):
    assert launch_cad_viewer(tmp_path / "nope.step") is False


def test_menu_and_interactive_subcommands(monkeypatch):
    called = []
    monkeypatch.setattr("cli.app.interactive_repl", lambda: called.append(True))
    assert main(["menu"]) == 0
    assert main(["interactive"]) == 0
    assert len(called) == 2


def test_render_menu_table_and_header():
    hdr = render_header()
    assert hdr is not None

    tbl = render_menu_table()
    assert tbl is not None
    assert len(tbl.columns) == 4


def test_interactive_menu_exit_on_q():
    test_console = Console()
    commands: list[list[str]] = []
    inputs = iter(["q"])
    test_console.input = lambda prompt="": next(inputs)  # type: ignore[assignment]

    run_interactive_menu(lambda cmd: commands.append(cmd) or 0, test_console)
    assert commands == []


def test_interactive_menu_direct_command_dispatch():
    test_console = Console()
    commands: list[list[str]] = []
    inputs = iter(["info", "", "q"])
    test_console.input = lambda prompt="": next(inputs)  # type: ignore[assignment]

    run_interactive_menu(lambda cmd: commands.append(cmd) or 0, test_console)
    assert commands == [["info"]]


def test_interactive_menu_chair_flow():
    test_console = Console()
    commands: list[list[str]] = []
    # Choice 1 (chair), Option 1 (standard), 'n' for CAD viewer, '' for Enter to return, 'q' to exit
    inputs = iter(["1", "1", "n", "", "q"])
    test_console.input = lambda prompt="": next(inputs)  # type: ignore[assignment]

    run_interactive_menu(lambda cmd: commands.append(cmd) or 0, test_console)
    assert commands == [["chair"]]


def test_interactive_menu_trace_flow():
    test_console = Console()
    commands: list[list[str]] = []
    # Choice 5 (trace), Option 1 (overview), Enter to continue, 'b' to back, Enter to return, 'q' to exit
    inputs = iter(["5", "1", "", "b", "", "q"])
    test_console.input = lambda prompt="": next(inputs)  # type: ignore[assignment]

    run_interactive_menu(lambda cmd: commands.append(cmd) or 0, test_console)
    assert commands == [["trace"]]


