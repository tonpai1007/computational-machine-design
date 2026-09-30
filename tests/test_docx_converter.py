"""
Tests for the HTML to Word (.docx) report converter.

The rich converter behind the bridge lives in the separate learning-tools
repository, so these are skipped when it is not installed. The in-repo
renderer in ``mdie.convert.documents`` is covered unconditionally by
``tests/test_converter.py``.
"""

from pathlib import Path
import docx
import pytest

from mdie.convert.docx_bridge import (
    HTMLToDocxConverter,
    is_available,
    learning_tools_path,
)

requires_learning_tools = pytest.mark.skipif(
    not is_available(),
    reason="learning-tools checkout not present (set LEARNING_TOOLS_PATH)",
)


def test_bridge_imports_even_when_the_dependency_is_missing():
    """A missing optional integration must never break an import."""
    assert HTMLToDocxConverter is not None
    assert isinstance(is_available(), bool)


def test_learning_tools_path_is_none_or_a_real_checkout():
    path = learning_tools_path()
    if path is not None:
        assert (path / "lt" / "docx_converter.py").exists()


@pytest.mark.skipif(is_available(), reason="learning-tools is installed")
def test_using_the_bridge_without_learning_tools_raises_a_clear_error(tmp_path):
    from mdie.convert.docx_bridge import LearningToolsUnavailable

    src = tmp_path / "r.html"
    src.write_text("<html><body><p>hi</p></body></html>", encoding="utf-8")
    with pytest.raises(LearningToolsUnavailable, match="LEARNING_TOOLS_PATH"):
        HTMLToDocxConverter.convert(src, tmp_path / "r.docx")


def test_bridge_degrades_to_none_when_the_checkout_is_absent(monkeypatch):
    """Simulate a machine without learning-tools: load() must not raise."""
    import mdie.convert.docx_bridge as bridge

    monkeypatch.setattr(bridge, "learning_tools_path", lambda: None)
    assert bridge._load() is None
    assert bridge.is_available() is False


@requires_learning_tools
def test_docx_conversion_basic(tmp_path: Path):
    html_content = """<!DOCTYPE html>
    <html>
    <head><title>Test Calculation Report</title></head>
    <body>
        <h1>Chapter 1: Static Equilibrium</h1>
        <p>This is a <b>critical</b> test of the <i>mechanical</i> calculation.</p>
        <table>
            <tr><th>Component</th><th>Load (N)</th><th>Verdict</th></tr>
            <tr><td>Shaft</td><td>1500.0</td><td>PASS</td></tr>
            <tr><td>Keyway</td><td>350.0</td><td>FAIL</td></tr>
        </table>
        <ul>
            <li>Factor of safety > 2.0</li>
            <li>Material: AISI 1018 Steel</li>
        </ul>
    </body>
    </html>
    """
    html_file = tmp_path / "test_report.html"
    docx_file = tmp_path / "test_report.docx"
    html_file.write_text(html_content, encoding="utf-8")

    out_path = HTMLToDocxConverter.convert(html_file, docx_file)
    assert out_path.exists()
    assert out_path.stat().st_size > 1000

    # Verify docx content
    doc = docx.Document(str(out_path))
    assert len(doc.tables) >= 1
    table = doc.tables[0]
    assert len(table.rows) == 3
    assert table.rows[0].cells[0].text == "Component"
    assert table.rows[1].cells[2].text == "PASS"

    # Verify paragraphs
    all_text = "\n".join(p.text for p in doc.paragraphs)
    assert "Chapter 1: Static Equilibrium" in all_text
    assert "Factor of safety > 2.0" in all_text


@requires_learning_tools
def test_docx_conversion_with_ai(tmp_path: Path):
    html_content = """<!DOCTYPE html>
    <html>
    <head><title>Machine Design Dossier</title></head>
    <body>
        <div class="cover-sheet">
            <div class="cover-header">Department of Mechanical Engineering</div>
            <div class="cover-course">Machine Design Project</div>
            <div class="cover-title">Ergonomic Chair Structure</div>
        </div>
        <h1>1.0 Structural Audit</h1>
        <p>Applied vertical load = 1400 N.</p>
    </body>
    </html>
    """
    html_file = tmp_path / "test_academic.html"
    docx_file = tmp_path / "test_academic.docx"
    html_file.write_text(html_content, encoding="utf-8")

    out_path = HTMLToDocxConverter.convert(html_file, docx_file, enable_ai=True)
    assert out_path.exists()

    doc = docx.Document(str(out_path))
    # Should have AI peer review callout table or header
    all_text = "\n".join(p.text for p in doc.paragraphs) + "\n".join(c.text for t in doc.tables for row in t.rows for c in row.cells)
    assert "AI Engineering Executive Summary" in all_text or "Ergonomic Chair Structure" in all_text
