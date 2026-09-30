"""Document conversion: Markdown / HTML / DOCX / PDF / TXT.

Everything is parsed into a small block model and then rendered to the target
format, so any source/target pair is a single function call rather than a
hand-written matrix of pairwise converters.

Block model
-----------
``("heading", level, text)`` ``("para", text)`` ``("bullet", text)``
``("numbered", text)`` ``("code", text)`` ``("table", rows)`` ``("quote", text)``
where ``rows`` is a list of lists of cell text.
"""

from __future__ import annotations

import html as html_mod
import re
from html.parser import HTMLParser
from pathlib import Path
from typing import Iterable

Block = tuple

# ---------------------------------------------------------------- parsing: md

_MD_HEADING = re.compile(r"^(#{1,6})\s+(.*)$")
_MD_BULLET = re.compile(r"^\s*[-*+]\s+(.*)$")
_MD_NUMBERED = re.compile(r"^\s*\d+[.)]\s+(.*)$")
_MD_TABLE_SEP = re.compile(r"^\s*\|?[\s:|-]+\|[\s:|-]*$")


def blocks_from_markdown(text: str) -> list[Block]:
    """Parse Markdown into blocks, honouring ATX headings, lists, code and tables."""
    lines = text.replace("\r\n", "\n").split("\n")
    blocks: list[Block] = []
    i = 0
    para: list[str] = []

    def flush() -> None:
        if para:
            joined = " ".join(para).strip()
            if joined:
                blocks.append(("para", joined))
            para.clear()

    while i < len(lines):
        line = lines[i]

        if line.strip().startswith("```"):
            flush()
            i += 1
            code: list[str] = []
            while i < len(lines) and not lines[i].strip().startswith("```"):
                code.append(lines[i])
                i += 1
            blocks.append(("code", "\n".join(code)))
            i += 1
            continue

        m = _MD_HEADING.match(line)
        if m:
            flush()
            blocks.append(("heading", len(m.group(1)), m.group(2).strip()))
            i += 1
            continue

        if line.lstrip().startswith("> "):
            flush()
            quote: list[str] = []
            while i < len(lines) and lines[i].lstrip().startswith("> "):
                quote.append(lines[i].lstrip()[2:])
                i += 1
            blocks.append(("quote", " ".join(quote).strip()))
            continue

        # Table: a header row followed by a separator row.
        if "|" in line and i + 1 < len(lines) and _MD_TABLE_SEP.match(lines[i + 1]):
            flush()
            rows = [_split_table_row(line)]
            i += 2
            while i < len(lines) and "|" in lines[i] and lines[i].strip():
                rows.append(_split_table_row(lines[i]))
                i += 1
            blocks.append(("table", rows))
            continue

        m = _MD_BULLET.match(line)
        if m:
            flush()
            blocks.append(("bullet", m.group(1).strip()))
            i += 1
            continue

        m = _MD_NUMBERED.match(line)
        if m:
            flush()
            blocks.append(("numbered", m.group(1).strip()))
            i += 1
            continue

        if not line.strip():
            flush()
            i += 1
            continue

        para.append(line.strip())
        i += 1

    flush()
    return blocks


def _split_table_row(line: str) -> list[str]:
    cells = line.strip().strip("|").split("|")
    return [c.strip() for c in cells]


# --------------------------------------------------------------- parsing: html

class _HtmlBlocks(HTMLParser):
    """Collect block-level text out of an HTML document."""

    def __init__(self) -> None:
        super().__init__(convert_charrefs=True)
        self.blocks: list[Block] = []
        self._tag: str | None = None
        self._level = 0
        self._buf: list[str] = []
        self._in_code = False
        self._list_tag: str | None = None
        self._table: list[list[str]] | None = None
        self._row: list[str] | None = None
        self._cell: list[str] | None = None
        self._in_cell = False
        self._skip = False

    # -- helpers
    def _flush(self, kind: str, extra=None) -> None:
        text = re.sub(r"\s+", " ", "".join(self._buf)).strip()
        self._buf.clear()
        if not text and kind != "table":
            return
        self.blocks.append((kind, *(extra or ()), text) if extra else (kind, text))

    # -- parser hooks
    def handle_starttag(self, tag, attrs):
        if tag in ("head", "style", "script"):
            self._skip = True
            return
        if tag in ("h1", "h2", "h3", "h4", "h5", "h6"):
            self._flush("para")
            self._tag, self._level = tag, int(tag[1])
        elif tag in ("p", "div", "blockquote"):
            if tag == "blockquote":
                self._flush("para")
                self._tag = "quote"
            else:
                self._flush("para")
                self._tag = tag
        elif tag == "li":
            self._flush("para")
            self._tag = "bullet" if self._list_tag == "ul" else "numbered"
        elif tag in ("ul", "ol"):
            self._flush("para")
            self._list_tag = tag
        elif tag == "pre":
            self._flush("para")
            self._in_code = True
            self._tag = "code"
        elif tag == "table":
            self._flush("para")
            self._table, self._row = [], []
        elif tag in ("td", "th"):
            self._in_cell, self._cell = True, []

    def handle_endtag(self, tag):
        if tag in ("head", "style", "script"):
            self._skip = False
            return
        if tag in ("h1", "h2", "h3", "h4", "h5", "h6"):
            self._flush("heading", (self._level,))
            self._tag = None
        elif tag == "blockquote":
            self._flush("quote")
            self._tag = None
        elif tag == "li":
            self._flush(self._tag or "bullet")
            self._tag = None
        elif tag in ("ul", "ol"):
            self._list_tag = None
        elif tag in ("p", "div"):
            self._flush("para")
            self._tag = None
        elif tag == "pre":
            self._in_code = False
            self._flush("code")
            self._tag = None
        elif tag == "tr":
            if self._table is not None and self._row:
                self._table.append(self._row)
            self._row = []
        elif tag in ("td", "th"):
            if self._cell is not None and self._row is not None:
                self._row.append(re.sub(r"\s+", " ", "".join(self._cell)).strip())
            self._in_cell = False
        elif tag == "table":
            self._flush("para")
            if self._table:
                width = max((len(r) for r in self._table), default=0)
                rows = [r + [""] * (width - len(r)) for r in self._table]
                self.blocks.append(("table", rows))
            self._table = None

    def handle_data(self, data):
        if self._skip:
            return
        if self._in_cell and self._cell is not None:
            self._cell.append(data)
            return
        self._buf.append(data)

    def close(self):
        super().close()
        self._flush("para")
        if self._table:
            self.blocks.append(("table", self._table))


def blocks_from_html(text: str) -> list[Block]:
    """Parse an HTML document into blocks."""
    parser = _HtmlBlocks()
    parser.feed(text)
    parser.close()
    return parser.blocks


# --------------------------------------------------------------- parsing: docx

def blocks_from_docx(path: Path) -> list[Block]:
    """Read a Word document back into blocks."""
    import docx as python_docx

    document = python_docx.Document(str(path))
    blocks: list[Block] = []
    for para in document.paragraphs:
        text = para.text.strip()
        if not text:
            continue
        style = (para.style.name or "").lower()
        if style.startswith("heading"):
            digits = "".join(ch for ch in style if ch.isdigit())
            blocks.append(("heading", int(digits) if digits else 1, text))
        elif style.startswith("list bullet"):
            blocks.append(("bullet", text))
        elif style.startswith("list number"):
            blocks.append(("numbered", text))
        elif style.startswith("quote"):
            blocks.append(("quote", text))
        else:
            blocks.append(("para", text))

    for table in document.tables:
        rows = [[c.text.strip() for c in row.cells] for row in table.rows]
        if rows:
            blocks.append(("table", rows))
    return blocks


# --------------------------------------------------------------- parsing: pdf

def blocks_from_pdf(path: Path) -> list[Block]:
    """Extract text from a PDF into paragraphs."""
    import fitz

    blocks: list[Block] = []
    with fitz.open(str(path)) as doc:
        for page in doc:
            for raw in page.get_text("blocks"):
                text = re.sub(r"\s+", " ", raw[4]).strip()
                if text:
                    blocks.append(("para", text))
    return blocks


# ------------------------------------------------------------------- renderers

_MD_ESCAPE = re.compile(r"([\\`*_\[\]])")

_CSS = """
body { font-family: Helvetica, Arial, sans-serif; font-size: 10.5pt; line-height: 1.45; }
h1 { font-size: 17pt; } h2 { font-size: 14pt; } h3 { font-size: 12pt; }
code, pre { font-family: Courier, monospace; font-size: 9pt; background: #f2f2f2; }
table { border-collapse: collapse; } td, th { border: 1px solid #999; padding: 3px 6px; }
"""


def blocks_to_markdown(blocks: Iterable[Block]) -> str:
    out: list[str] = []
    for block in blocks:
        kind = block[0]
        if kind == "heading":
            out += [f"{'#' * block[1]} {block[2]}", ""]
        elif kind == "para":
            out += [block[1], ""]
        elif kind == "bullet":
            out.append(f"- {block[1]}")
        elif kind == "numbered":
            out.append(f"1. {block[1]}")
        elif kind == "quote":
            out.append(f"> {block[1]}")
        elif kind == "code":
            out += ["```", block[1], "```", ""]
        elif kind == "table":
            rows = block[1]
            if not rows:
                continue
            width = max(len(r) for r in rows)
            rows = [r + [""] * (width - len(r)) for r in rows]
            out.append("| " + " | ".join(rows[0]) + " |")
            out.append("|" + "|".join(["---"] * width) + "|")
            for row in rows[1:]:
                out.append("| " + " | ".join(row) + " |")
            out.append("")
    return "\n".join(out).strip() + "\n"


def blocks_to_html(blocks: Iterable[Block], title: str = "MDIE Document") -> str:
    body: list[str] = []
    for block in blocks:
        kind = block[0]
        if kind == "heading":
            level = min(block[1] + 1, 6)
            body.append(f"<h{level}>{html_mod.escape(block[2])}</h{level}>")
        elif kind == "para":
            body.append(f"<p>{html_mod.escape(block[1])}</p>")
        elif kind == "bullet":
            body.append(f"<ul><li>{html_mod.escape(block[1])}</li></ul>")
        elif kind == "numbered":
            body.append(f"<ol><li>{html_mod.escape(block[1])}</li></ol>")
        elif kind == "quote":
            body.append(f"<blockquote>{html_mod.escape(block[1])}</blockquote>")
        elif kind == "code":
            body.append(f"<pre>{html_mod.escape(block[1])}</pre>")
        elif kind == "table":
            rows = block[1]
            cells = "".join("<tr>" + "".join(f"<td>{html_mod.escape(c)}</td>" for c in r) + "</tr>"
                            for r in rows)
            body.append(f"<table>{cells}</table>")
    return (f"<!DOCTYPE html><html><head><meta charset='utf-8'>"
            f"<title>{html_mod.escape(title)}</title></head>"
            f"<body>{''.join(body)}</body></html>")


def blocks_to_text(blocks: Iterable[Block]) -> str:
    out: list[str] = []
    for block in blocks:
        kind = block[0]
        if kind == "heading":
            out += ["", block[2].upper(), "=" * len(block[2]), ""]
        elif kind == "bullet":
            out.append(f"  * {block[1]}")
        elif kind == "numbered":
            out.append(f"  {block[1]}")
        elif kind == "para":
            out += [block[1], ""]
        elif kind == "quote":
            out += [f"  | {block[1]}"]
        elif kind == "code":
            out += ["", *[f"  {ln}" for ln in block[1].splitlines()], ""]
        elif kind == "table":
            for row in block[1]:
                out.append("  " + " | ".join(row))
            out.append("")
    return "\n".join(out).strip() + "\n"


def blocks_to_docx(blocks: Iterable[Block], path: Path, title: str = "MDIE Document") -> Path:
    import docx as python_docx

    document = python_docx.Document()
    for block in blocks:
        kind = block[0]
        if kind == "heading":
            document.add_heading(block[2], level=min(block[1], 9))
        elif kind == "para":
            document.add_paragraph(block[1])
        elif kind == "bullet":
            document.add_paragraph(block[1], style="List Bullet")
        elif kind == "numbered":
            document.add_paragraph(block[1], style="List Number")
        elif kind == "quote":
            document.add_paragraph(block[1], style="Intense Quote")
        elif kind == "code":
            para = document.add_paragraph()
            para.add_run(block[1]).font.name = "Courier New"
        elif kind == "table":
            rows = block[1]
            if not rows:
                continue
            table = document.add_table(rows=len(rows), cols=max(len(r) for r in rows))
            table.style = "Table Grid"
            for r_idx, row in enumerate(rows):
                for c_idx, cell in enumerate(row):
                    table.cell(r_idx, c_idx).text = cell
    path.parent.mkdir(parents=True, exist_ok=True)
    document.save(str(path))
    return path


def blocks_to_pdf(blocks: Iterable[Block], path: Path, title: str = "MDIE Document") -> Path:
    """Lay the blocks out as a paginated PDF using PyMuPDF's HTML story engine."""
    import fitz

    markup = blocks_to_html(blocks, title)
    path.parent.mkdir(parents=True, exist_ok=True)

    mediabox = fitz.paper_rect("a4")
    story = fitz.Story(html=markup, user_css=_CSS, archive=fitz.Archive("."))
    writer = fitz.DocumentWriter(str(path))

    remaining = 1
    guard = 0
    while remaining and guard < 500:
        guard += 1
        device = writer.begin_page(mediabox)
        frame = fitz.Rect(50, 50, mediabox.width - 50, mediabox.height - 50)
        remaining, _ = story.place(frame)
        story.draw(device)
        writer.end_page()

    writer.close()
    return path


# ------------------------------------------------------------- source routing

_READERS = {
    "md": blocks_from_markdown,
    "html": blocks_from_html,
    "docx": blocks_from_docx,
    "pdf": blocks_from_pdf,
    "txt": None,
}


def blocks_from_source(path: Path, source_format: str) -> list[Block]:
    """Dispatch a source document onto the right block parser."""
    if source_format == "txt":
        return blocks_from_markdown(path.read_text(encoding="utf-8", errors="replace"))
    reader = _READERS.get(source_format)
    if reader is None:
        raise ValueError(f"cannot read '{source_format}' documents")
    if source_format in ("docx", "pdf"):
        return reader(path)
    return reader(path.read_text(encoding="utf-8", errors="replace"))


_RENDERERS = {
    "md": blocks_to_markdown,
    "html": blocks_to_html,
    "txt": blocks_to_text,
    "docx": blocks_to_docx,
    "pdf": blocks_to_pdf,
}